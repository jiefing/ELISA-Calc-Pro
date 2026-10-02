# -*- coding: utf-8 -*-
"""
核心逻辑中枢 (Controller)
协调 Model 层 (Plate/Well/WellType) 与 View 层 (MainWindow/PreviewDialog) 之间的动作交互。
集成多工作表 Excel 导出功能。
"""

import os
import logging
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from typing import Optional, List, Dict

from PyQt5.QtWidgets import QFileDialog, QMessageBox, QWidget

from models.data_structures import Plate, WellType, Well
from strategies import ALL_STRATEGIES
from strategies.base_strategy import StrategyError

logger = logging.getLogger("AppController")

class AppController:
    """
    主控制器中枢，调度 View 和 Model
    """
    def __init__(self, main_window=None):
        self.plate: Optional[Plate] = None
        self.current_strategy_index: int = 0
        self._main_window = main_window

        # 保存策略实例列表
        self._strategies = ALL_STRATEGIES

    def set_main_window(self, main_window):
        """绑定主窗口引用"""
        self._main_window = main_window

    def get_strategy_names(self) -> List[str]:
        """返回所有实验策略的中文名称"""
        return [s.name for s in self._strategies]

    def get_current_strategy(self):
        """获取当前激活的实验策略实例"""
        return self._strategies[self.current_strategy_index]

    # ─────────────────────────────────────────────────
    # 导入文件逻辑
    # ─────────────────────────────────────────────────

    def import_file(self):
        """从对话框选择并拉起 PreviewDialog 确认数据"""
        if self._main_window is None:
            return

        filepath, _ = QFileDialog.getOpenFileName(
            self._main_window,
            "打开酶标仪数据文件",
            "",
            "数据文件 (*.csv *.xlsx *.xls *.txt);;"
            "Excel 文件 (*.xlsx *.xls);;"
            "CSV 逗号分割文件 (*.csv);;"
            "TXT 文本文件 (*.txt)"
        )

        if not filepath:
            self._main_window.status_bar.showMessage("已取消导入。")
            return

        logger.info(f"拉起手动框选对话框加载文件: {filepath}")

        try:
            # 引入重定义的手动位置框选对话框
            from views.preview_dialog import PreviewDialog
            dialog = PreviewDialog(filepath, self._main_window)

            if dialog.exec_() == PreviewDialog.Accepted:
                plate = dialog.get_imported_plate()
                self.plate = plate
                self._main_window.load_plate_data(plate)
                self._main_window.status_bar.showMessage(
                    f"导入成功: {os.path.basename(filepath)}"
                )
            else:
                self._main_window.status_bar.showMessage("已取消确认导入。")
        except Exception as e:
            logger.error(f"拉起确认预览框崩溃异常：{e}", exc_info=True)
            QMessageBox.critical(
                self._main_window, "导入失败",
                f"拉起手动框选预览界面失败，崩溃信息：\n{e}"
            )

    # ─────────────────────────────────────────────────
    # 运算执行逻辑
    # ─────────────────────────────────────────────────

    def run_calculation(self, strategy_index: int, params: dict):
        """驱动选定策略进行数据计算"""
        if self.plate is None:
            raise StrategyError("尚未导入任何有效酶标仪孔板数据！")

        self.current_strategy_index = strategy_index
        strategy = self.get_current_strategy()

        # 检查是否进行了分类标记
        labeled_wells = [w for w in self.plate.all_wells() if w.well_type != WellType.UNUSED]
        if not labeled_wells:
            raise StrategyError(
                "请先在孔板上标记数据类型！\n"
                "（用鼠标左键拖拽选中微孔，右键设置类型，如 Standard 或 Sample）"
            )

        # 还原计算状态
        for well in self.plate.all_wells():
            well.calc_result = None

        # 驱动核心计算 (任何计算基于 Well.corrected_value)
        self.plate = strategy.calculate(self.plate, params)

    # ─────────────────────────────────────────────────
    # Excel 报表写入与导出 (合并原 ExportManager 模块)
    # ─────────────────────────────────────────────────

    def export_report(self):
        """弹出保存框，导出 Excel 报表"""
        if self.plate is None:
            QMessageBox.warning(self._main_window, "警告", "暂无可用数据，无法执行导出报表！")
            return

        strategy = self.get_current_strategy()

        # 默认保存名称
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_filename = f"酶标仪结果_策略-{strategy_index_to_str(self.current_strategy_index)}_{timestamp}.xlsx"

        filepath, _ = QFileDialog.getSaveFileName(
            self._main_window,
            "保存分析报表",
            default_filename,
            "Excel 工作表 (*.xlsx)"
        )

        if not filepath:
            return

        if not filepath.lower().endswith(".xlsx"):
            filepath += ".xlsx"

        try:
            # 执行 Excel 多工作表导出
            params = self._main_window.param_panel.get_current_params()
            unit = params.get("unit", "ng/mL")
            self._write_excel_report(filepath, strategy.name, unit, params)

            QMessageBox.information(
                self._main_window, "导出成功",
                f"报表已成功保存并写出到：\n{filepath}"
            )
        except Exception as e:
            logger.error(f"导出 Excel 报表失败：{e}", exc_info=True)
            QMessageBox.critical(
                self._main_window, "导出错误",
                f"写出 Excel 报表发生致命故障：\n{e}"
            )

    def _write_excel_report(self, filepath: str, strategy_name: str, unit: str, params: dict):
        """执行具体的 openpyxl 多工作表写入"""
        wb = openpyxl.Workbook()

        # 统一表头填充（暗深蓝紫色）
        header_fill = PatternFill("solid", fgColor="2D2B55")
        subheader_fill = PatternFill("solid", fgColor="3D3B65")

        # 统一微孔基色填充
        type_fills = {
            WellType.BLANK:    PatternFill("solid", fgColor="F2F2F2"),
            WellType.STANDARD: PatternFill("solid", fgColor="FFF2F2"),
            WellType.SAMPLE:   PatternFill("solid", fgColor="F0F4FF"),
            WellType.CONTROL:  PatternFill("solid", fgColor="F0FFF0"),
            WellType.CUSTOM:   PatternFill("solid", fgColor="FFFAF0"),
            WellType.UNUSED:   PatternFill("solid", fgColor="FFFFFF"),
        }

        # -------------------------------------------------------------
        # Sheet 1: OD 原始数据热图
        # -------------------------------------------------------------
        ws_od = wb.active
        ws_od.title = "原始OD值"

        thin_side = Side(style="thin", color="AAAAAA")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        center_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)

        # 标题栏
        ws_od.merge_cells("A1:M1")
        title_cell = ws_od["A1"]
        title_cell.value = "酶标仪原始 OD 读数矩阵"
        title_cell.font = Font(name="Microsoft YaHei", bold=True, size=12, color="FFFFFF")
        title_cell.fill = header_fill
        title_cell.alignment = center_align

        # 列首
        ws_od.cell(row=2, column=1, value="").fill = subheader_fill
        for c in range(1, self.plate.cols + 1):
            cell = ws_od.cell(row=2, column=c + 1, value=str(c))
            cell.font = header_font
            cell.fill = subheader_fill
            cell.alignment = center_align
            cell.border = cell_border
            ws_od.column_dimensions[get_column_letter(c + 1)].width = 10
        ws_od.column_dimensions["A"].width = 6

        # 行首及数值
        for r_idx, row_label in enumerate(self.plate.row_labels):
            row = 3 + r_idx
            ws_od.row_dimensions[row].height = 20

            row_head = ws_od.cell(row=row, column=1, value=row_label)
            row_head.font = header_font
            row_head.fill = subheader_fill
            row_head.alignment = center_align
            row_head.border = cell_border

            for c_idx in range(1, self.plate.cols + 1):
                well = self.plate.get_well(row_label, c_idx)
                cell = ws_od.cell(row=row, column=c_idx + 1)
                cell.alignment = center_align
                cell.border = cell_border

                if well and well.raw_value is not None:
                    cell.value = well.raw_value
                    cell.number_format = "0.000"
                else:
                    cell.value = ""

                # 孔类型着色
                if well:
                    fill = type_fills.get(well.well_type)
                    if fill:
                        cell.fill = fill

        # -------------------------------------------------------------
        # Sheet 2: 计算结果明细表
        # -------------------------------------------------------------
        ws_detail = wb.create_sheet("计算结果明细")
        ws_detail.merge_cells("A1:E1")

        d_title = ws_detail["A1"]
        d_title.value = f"实验计算结果明细表 --- 策略: {strategy_name}"
        d_title.font = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
        d_title.fill = header_fill
        d_title.alignment = center_align

        # 表头
        headers = ["孔位", "孔类型", "原始OD", "计算结果", "备注"]
        for idx, h in enumerate(headers, 1):
            cell = ws_detail.cell(row=2, column=idx, value=h)
            cell.font = header_font
            cell.fill = subheader_fill
            cell.alignment = center_align
            cell.border = cell_border

        # 收集非 unused 孔写入数据行
        data_wells = [w for w in self.plate.all_wells() if w.well_type != WellType.UNUSED]
        # 按 A1-H12 空间排序
        data_wells.sort(key=lambda w: (w.row, w.col))

        type_names = {
            WellType.BLANK: "Blank 空白",
            WellType.STANDARD: "Standard 标曲",
            WellType.SAMPLE: "Sample 样品",
            WellType.CONTROL: "Control 对照",
            WellType.CUSTOM: "Custom 自定义",
        }

        for idx, well in enumerate(data_wells, 3):
            ws_detail.cell(row=idx, column=1, value=well.label).alignment = center_align
            ws_detail.cell(row=idx, column=2, value=type_names.get(well.well_type, "Unused")).alignment = center_align

            od_cell = ws_detail.cell(row=idx, column=3, value=well.raw_value)
            od_cell.alignment = center_align
            od_cell.number_format = "0.000"

            res_cell = ws_detail.cell(row=idx, column=4)
            res_cell.alignment = center_align
            if well.calc_result is not None:
                if isinstance(well.calc_result, float):
                    res_cell.value = well.calc_result
                    res_cell.number_format = "0.000"
                else:
                    res_cell.value = str(well.calc_result) # 拦截的 >Max 等字符
            else:
                res_cell.value = ""

            ws_detail.cell(row=idx, column=5, value="").alignment = center_align

            # 单独上色
            fill = type_fills.get(well.well_type)
            if fill:
                ws_detail.cell(row=idx, column=2).fill = fill

            for c in range(1, 6):
                ws_detail.cell(row=idx, column=c).border = cell_border

        for col, w in zip("ABCDE", [10, 15, 14, 18, 16]):
            ws_detail.column_dimensions[col].width = w

        # -------------------------------------------------------------
        # Sheet 3: 计算结果矩阵图
        # -------------------------------------------------------------
        ws_res_map = wb.create_sheet("计算结果矩阵")
        ws_res_map.merge_cells("A1:M1")

        m_title = ws_res_map["A1"]
        m_title.value = f"计算结果空间分布图 ({unit})"
        m_title.font = Font(name="Microsoft YaHei", bold=True, size=12, color="FFFFFF")
        m_title.fill = header_fill
        m_title.alignment = center_align

        # 列头
        ws_res_map.cell(row=2, column=1, value="").fill = subheader_fill
        for c in range(1, self.plate.cols + 1):
            cell = ws_res_map.cell(row=2, column=c + 1, value=str(c))
            cell.font = header_font
            cell.fill = subheader_fill
            cell.alignment = center_align
            cell.border = cell_border
            ws_res_map.column_dimensions[get_column_letter(c + 1)].width = 11

        ws_res_map.column_dimensions["A"].width = 6

        # 数据行
        for r_idx, row_label in enumerate(self.plate.row_labels):
            row = 3 + r_idx
            ws_res_map.row_dimensions[row].height = 20

            row_head = ws_res_map.cell(row=row, column=1, value=row_label)
            row_head.font = header_font
            row_head.fill = subheader_fill
            row_head.alignment = center_align
            row_head.border = cell_border

            for c_idx in range(1, self.plate.cols + 1):
                well = self.plate.get_well(row_label, c_idx)
                cell = ws_res_map.cell(row=row, column=c_idx + 1)
                cell.alignment = center_align
                cell.border = cell_border

                if well and well.calc_result is not None:
                    if isinstance(well.calc_result, float):
                        cell.value = well.calc_result
                        cell.number_format = "0.000"
                    else:
                        cell.value = str(well.calc_result)
                else:
                    cell.value = ""

                if well:
                    fill = type_fills.get(well.well_type)
                    if fill:
                        cell.fill = fill

        # -------------------------------------------------------------
        # Sheet 4: 样品计算汇总 (Sample Summary)
        # -------------------------------------------------------------
        ws_sum = wb.create_sheet("样品计算汇总")
        ws_sum.merge_cells("A1:E1")

        sum_title = ws_sum["A1"]
        sum_title.value = f"样品复孔计算汇总明细 ({unit}) --- 策略: {strategy_name}"
        sum_title.font = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
        sum_title.fill = header_fill
        sum_title.alignment = center_align

        # 表头
        sum_headers = ["样本编号", "包含孔位", "OD均值", "最终结果均值", "复孔CV%"]
        for idx, h in enumerate(sum_headers, 1):
            cell = ws_sum.cell(row=2, column=idx, value=h)
            cell.font = header_font
            cell.fill = subheader_fill
            cell.alignment = center_align
            cell.border = cell_border

        # 读取 sample_summary 统计
        sample_sum = self.plate.statistics.get("sample_summary", {})

        # 排序：按照 S1, S2, S3... 的数字大小顺序进行排序
        def get_group_sort_key(g_id: str):
            if g_id.startswith("S"):
                try:
                    return (0, int(g_id[1:]))
                except ValueError:
                    pass
            return (1, g_id)

        sorted_groups = sorted(sample_sum.keys(), key=get_group_sort_key)

        for idx, g_id in enumerate(sorted_groups, 3):
            info = sample_sum[g_id]

            ws_sum.cell(row=idx, column=1, value=g_id).alignment = center_align
            ws_sum.cell(row=idx, column=2, value=", ".join(info["wells"])).alignment = center_align

            od_c = ws_sum.cell(row=idx, column=3, value=info["od_mean"])
            od_c.alignment = center_align
            od_c.number_format = "0.000"

            calc_c = ws_sum.cell(row=idx, column=4, value=info["calc_mean"])
            calc_c.alignment = center_align
            calc_c.number_format = "0.000"

            cv_c = ws_sum.cell(row=idx, column=5, value=f"{info['calc_cv']:.2f}%" if info["calc_cv"] > 0 else "—")
            cv_c.alignment = center_align

            # 样品底色
            ws_sum.cell(row=idx, column=1).fill = type_fills[WellType.SAMPLE]

            for c in range(1, 6):
                ws_sum.cell(row=idx, column=c).border = cell_border

        for col, w in zip("ABCDE", [12, 25, 14, 18, 14]):
            ws_sum.column_dimensions[col].width = w

        # -------------------------------------------------------------
        # Sheet 5: 实验报告汇总与信息 (原 Sheet 4)
        # -------------------------------------------------------------
        ws_info = wb.create_sheet("报告摘要")
        ws_info.column_dimensions["A"].width = 22
        ws_info.column_dimensions["B"].width = 50

        # 元数据收集
        info_rows = [
            ("软件名称", "ELISA-Calc-Pro"),
            ("分析作者", "ELISA-Calc-Pro contributors"),
            ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("数据来源", self.plate.source_file or "未知路径"),
            ("孔板规格", f"{self.plate.plate_size} 孔板"),
            ("计算实验类型", strategy_name),
            ("计算数据单位", unit),
        ]

        # 4PL 的拟合参数
        if self.current_strategy_index == 1:
            popt = self.plate.statistics.get("popt")
            r_sq = self.plate.statistics.get("r_squared", 0.0)
            if popt:
                info_rows.append(("4PL 下限渐近值 (A)", f"{popt[0]:.6f}"))
                info_rows.append(("4PL 斜率因子 (B)", f"{popt[1]:.6f}"))
                info_rows.append(("4PL 效应浓度值 (C)", f"{popt[2]:.6f}"))
                info_rows.append(("4PL 上限渐近值 (D)", f"{popt[3]:.6f}"))
                info_rows.append(("4PL 决定系数 (R²)", f"{r_sq:.6f}"))

            cv_warns = self.plate.statistics.get("cv_warnings", [])
            if cv_warns:
                info_rows.append(("标曲质控报警 (CV%)", "; ".join(cv_warns)))

        # CCK-8
        elif self.current_strategy_index == 2:
            warns = self.plate.statistics.get("warnings", [])
            if warns:
                info_rows.append(("系统降级提醒", "; ".join(warns)))

        ws_info.merge_cells("A1:B1")
        title_c = ws_info["A1"]
        title_c.value = "📋 实验分析报表摘要"
        title_c.font = Font(name="Microsoft YaHei", bold=True, size=13, color="FFFFFF")
        title_c.fill = header_fill
        title_c.alignment = center_align

        for r_idx, (key, val) in enumerate(info_rows, 2):
            ws_info.row_dimensions[r_idx].height = 22

            k_cell = ws_info.cell(row=r_idx, column=1, value=key)
            k_cell.font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
            k_cell.fill = subheader_fill
            k_cell.alignment = Alignment(horizontal="right", vertical="center")
            k_cell.border = cell_border

            v_cell = ws_info.cell(row=r_idx, column=2, value=str(val))
            v_cell.font = Font(name="Microsoft YaHei", size=10)
            v_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            v_cell.border = cell_border

        wb.save(filepath)


def strategy_index_to_str(idx: int) -> str:
    modes = {0: "扣除空白", 1: "4PL标曲", 2: "CCK8活力"}
    return modes.get(idx, "未知")
